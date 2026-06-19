from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = r"C:\Loco\LocalProjects_01\GoMassive\next-mef\public\images\mef_plain_logo_horz.png"
RED = RGBColor(0xE5,0x00,0x00); BLACK = RGBColor(0,0,0); GREY = RGBColor(0x66,0x66,0x66)
doc = Document("reference-default.docx")

def setfont(style, name=None, size=None, bold=None, color=None):
    f = style.font
    if name:
        f.name = name
        rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a), name)
    if size is not None: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color

def bottom_border(style, color="E50000", sz=12):
    ppr = style.element.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),str(sz)); b.set(qn('w:space'),'4'); b.set(qn('w:color'),color)
    pbdr.append(b); ppr.append(pbdr)

# Body + structural styles
setfont(doc.styles['Normal'], "Inter", 10.5, False, BLACK)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
setfont(doc.styles['Title'], "Oswald", 30, True, BLACK)
setfont(doc.styles['Subtitle'], "Oswald", 15, False, RED)
setfont(doc.styles['Heading 1'], "Oswald", 17, True, BLACK); bottom_border(doc.styles['Heading 1'])
setfont(doc.styles['Heading 2'], "Oswald", 13.5, True, BLACK)
setfont(doc.styles['Heading 3'], "Oswald", 11.5, True, RED)
for h in ('Heading 1','Heading 2','Heading 3'):
    pf = doc.styles[h].paragraph_format; pf.space_before = Pt(12); pf.space_after = Pt(4); pf.keep_with_next = True
setfont(doc.styles['Block Text'], "Inter", 11.5, None, GREY)  # blockquote / epigraph

sec = doc.sections[0]
sec.different_first_page_header_footer = True  # clean cover

# ---- running header (logo, right aligned) ----
hp = sec.header.paragraphs[0]; hp.text = ""; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture(LOGO, width=Inches(1.25))
sec.header.is_linked_to_previous = False

# ---- footer (label left, page number right) ----
fp = sec.footer.paragraphs[0]; fp.text = ""
fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
r1 = fp.add_run("Massive Earth Foundation  ·  MEFCAP — Draft 01"); r1.font.name="Inter"; r1.font.size=Pt(8); r1.font.color.rgb=GREY
r2 = fp.add_run("\tPage "); r2.font.name="Inter"; r2.font.size=Pt(8); r2.font.color.rgb=GREY
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),' PAGE ')
fr = OxmlElement('w:r'); ft = OxmlElement('w:t'); ft.text="1"; fr.append(ft); fld.append(fr)
fp._p.append(fld)
sec.footer.is_linked_to_previous = False

doc.save("reference.docx")
print("reference.docx built")
