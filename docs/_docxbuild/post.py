from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = r"C:\Loco\LocalProjects_01\GoMassive\next-mef\public\images\mef_plain_logo_horz.png"
RED = RGBColor(0xE5,0x00,0x00); BLACK = RGBColor(0,0,0); GREY = RGBColor(0x55,0x55,0x55)
doc = Document("content.docx")
body = doc.element.body
first = body[0]

def newp(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    first.addprevious(p._p)   # move to front, in call order
    return p

def run(p, text, font="Inter", size=11, bold=False, italic=False, color=BLACK):
    r = p.add_run(text); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color
    return r

# logo
p = newp(space_before=72, space_after=10); p.add_run().add_picture(LOGO, width=Inches(3.6))
# title block
run(newp(space_before=18, space_after=2), "Massive Earth Foundation", "Oswald", 26, True, color=BLACK)
run(newp(space_after=2), "Climate Action Portal (MEFCAP)", "Oswald", 26, True, color=RED)
run(newp(space_after=18), "A Gateway into Climate", "Oswald", 15, color=BLACK)
run(newp(space_after=30), "Activating the Community Layer of Massive Earth Foundation", "Inter", 12, italic=True, color=GREY)
# epigraph
run(newp(space_after=2), "“The greatest threat to our planet is the belief that somebody else will save it.”", "Inter", 11.5, italic=True, color=GREY)
run(newp(space_after=40), "— Robert Swan", "Inter", 10, color=GREY)
# meta
run(newp(space_after=2), "Author: Shaurya   ·   Date: 16 June 2026   ·   For: MEF Leadership", "Inter", 10, color=BLACK)
run(newp(space_after=0), "Amazon 6-Pager  ·  Draft 01", "Inter", 9, color=GREY)
# page break
pb = newp(space_after=0); pb.add_run().add_break(WD_BREAK.PAGE)

# enforce clean first page (cover): titlePg in sectPr
sectPr = body.find(qn('w:sectPr'))
if sectPr is not None and sectPr.find(qn('w:titlePg')) is None:
    sectPr.append(OxmlElement('w:titlePg'))

out = r"C:\Loco\LocalProjects_01\GoMassive\next-mef\docs\MEFCAP-Draft01.docx"
doc.save(out)
print("saved", out)
